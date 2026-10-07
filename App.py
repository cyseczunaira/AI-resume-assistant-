"""
AI Resume ATS Checker
---------------------
Upload a resume (PDF, DOCX or TXT), optionally paste a job description, and get:
  * an overall ATS score (0-100) with a breakdown by category
  * strengths and weaknesses
  * missing keywords
  * prioritised, actionable improvements
  * before/after rewrites of weak bullet points

UI: Streamlit    |    AI: Google Gemini Flash (via the google-genai SDK)
"""

import io
import os
from typing import List, Optional

import streamlit as st
from docx import Document
from google import genai
from google.genai import types
from pydantic import BaseModel, Field
from pypdf import PdfReader

# --------------------------------------------------------------------------- #
# Config
# --------------------------------------------------------------------------- #
DEFAULT_MODEL = os.getenv("GEMINI_MODEL", "gemini-2.5-flash")
MAX_FILE_MB = 5
MAX_RESUME_CHARS = 20_000  # keeps prompts small and fast
MIN_RESUME_CHARS = 150  # below this we assume the PDF is scanned / empty

SYSTEM_INSTRUCTION = """You are an expert technical recruiter and ATS (Applicant Tracking System) specialist.
You evaluate resumes the way modern ATS software and recruiters do: keyword match,
section structure, parseability, clarity, and measurable impact.

Rules:
- Be honest and specific. Do not inflate scores. A typical average resume scores 50-70.
- Only use information that is actually in the resume text. Never invent experience,
  employers, numbers or skills. In rewrites, use placeholders like [X%] or [number]
  where a metric is missing instead of making one up.
- If a job description is provided, judge keyword match against it. If not, judge
  against general best practice for the role the resume appears to target.
- Scores are integers from 0 to 100.
"""


# --------------------------------------------------------------------------- #
# Output schema (Gemini returns JSON that matches this)
# --------------------------------------------------------------------------- #
class CategoryScores(BaseModel):
    formatting_and_structure: int = Field(description="0-100: clear sections, ATS-friendly layout, consistent dates")
    keywords_and_skills: int = Field(description="0-100: relevant hard skills and keywords")
    experience_and_impact: int = Field(description="0-100: achievements, action verbs, measurable results")
    education_and_certifications: int = Field(description="0-100: relevant education and credentials")
    readability_and_length: int = Field(description="0-100: concise, scannable, free of errors")


class Improvement(BaseModel):
    priority: str = Field(description="One of: High, Medium, Low")
    area: str = Field(description="Short label, e.g. 'Summary', 'Skills', 'Experience'")
    issue: str = Field(description="What is wrong or missing")
    fix: str = Field(description="Concrete action the candidate should take")


class BulletRewrite(BaseModel):
    original: str = Field(description="A weak bullet taken from the resume")
    improved: str = Field(description="A stronger version. Use [X] placeholders for unknown metrics")


class ResumeAnalysis(BaseModel):
    overall_score: int = Field(description="0-100 overall ATS score")
    summary: str = Field(description="2-3 sentence overall assessment")
    category_scores: CategoryScores
    strengths: List[str]
    weaknesses: List[str]
    found_keywords: List[str] = Field(description="Relevant keywords already present")
    missing_keywords: List[str] = Field(description="Important keywords that are missing")
    improvements: List[Improvement]
    bullet_rewrites: List[BulletRewrite]


# --------------------------------------------------------------------------- #
# File parsing
# --------------------------------------------------------------------------- #
def extract_text_from_pdf(data: bytes) -> str:
    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception:
            raise ValueError("This PDF is password-protected. Please upload an unlocked copy.")
    pages = [(page.extract_text() or "") for page in reader.pages]
    return "\n".join(pages)


def extract_text_from_docx(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    # Many resumes keep skills/experience inside tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text.strip())
    return "\n".join(parts)


def extract_resume_text(filename: str, data: bytes) -> str:
    """Return clean text from an uploaded resume. Raises ValueError with a friendly message."""
    name = filename.lower()
    try:
        if name.endswith(".pdf"):
            text = extract_text_from_pdf(data)
        elif name.endswith(".docx"):
            text = extract_text_from_docx(data)
        elif name.endswith(".txt"):
            text = data.decode("utf-8", errors="ignore")
        else:
            raise ValueError("Unsupported file type. Please upload a PDF, DOCX or TXT file.")
    except ValueError:
        raise
    except Exception as exc:  # corrupted files etc.
        raise ValueError(f"Could not read this file ({type(exc).__name__}). Is it corrupted?") from exc

    text = "\n".join(line.rstrip() for line in text.splitlines()).strip()
    if len(text) < MIN_RESUME_CHARS:
        raise ValueError(
            "Almost no text could be extracted. If your resume is a scanned image or "
            "was exported as pictures, ATS software can't read it either, which is a "
            "major problem in itself. Export a text-based PDF or DOCX and try again."
        )
    return text[:MAX_RESUME_CHARS]


# --------------------------------------------------------------------------- #
# Gemini call
# --------------------------------------------------------------------------- #
def build_prompt(resume_text: str, job_description: Optional[str]) -> str:
    jd = job_description.strip() if job_description else ""
    if jd:
        jd_block = f"JOB DESCRIPTION:\n\"\"\"\n{jd[:8000]}\n\"\"\"\n"
        task = "Score this resume against the job description above."
    else:
        jd_block = "JOB DESCRIPTION: (none provided)\n"
        task = "No job description was given, so score against general ATS best practice for the role this resume targets."
    return (
        f"{task}\n\n{jd_block}\n"
        f"RESUME TEXT:\n\"\"\"\n{resume_text}\n\"\"\"\n\n"
        "Give 5-8 prioritised improvements and 3-5 bullet rewrites."
    )


def _clamp(n: int) -> int:
    return max(0, min(100, int(n)))


def normalise(result: ResumeAnalysis) -> ResumeAnalysis:
    """Guard against out-of-range numbers from the model."""
    result.overall_score = _clamp(result.overall_score)
    for field in CategoryScores.model_fields:
        setattr(result.category_scores, field, _clamp(getattr(result.category_scores, field)))
    return result


def analyze_resume(
    api_key: str,
    resume_text: str,
    job_description: Optional[str] = None,
    model: str = DEFAULT_MODEL,
    client: Optional[genai.Client] = None,
) -> ResumeAnalysis:
    client = client or genai.Client(api_key=api_key)
    response = client.models.generate_content(
        model=model,
        contents=build_prompt(resume_text, job_description),
        config=types.GenerateContentConfig(
            system_instruction=SYSTEM_INSTRUCTION,
            response_mime_type="application/json",
            response_schema=ResumeAnalysis,
            temperature=0.3,
        ),
    )
    parsed = getattr(response, "parsed", None)
    if isinstance(parsed, ResumeAnalysis):
        return normalise(parsed)
    if not getattr(response, "text", None):
        raise RuntimeError("The model returned an empty response (it may have been blocked). Please try again.")
    return normalise(ResumeAnalysis.model_validate_json(response.text))


def friendly_error(exc: Exception) -> str:
    msg = str(exc)
    low = msg.lower()
    if "api key" in low or "api_key" in low or "permission" in low or "401" in low or "403" in low:
        return "The API key was rejected. Check that your Gemini API key is correct."
    if "429" in low or "quota" in low or "resource_exhausted" in low:
        return "Rate limit or quota reached. Wait a minute and try again, or check your Gemini plan."
    if "404" in low or "not found" in low:
        return f"The model name wasn't found. Try another model in the sidebar. ({msg[:150]})"
    if "503" in low or "unavailable" in low or "overloaded" in low:
        return "The Gemini service is busy right now. Please try again in a moment."
    return f"Something went wrong: {msg[:300]}"


# --------------------------------------------------------------------------- #
# UI helpers
# --------------------------------------------------------------------------- #
def score_label(score: int) -> str:
    if score >= 80:
        return "Excellent"
    if score >= 65:
        return "Good"
    if score >= 50:
        return "Needs work"
    return "Poor"


def build_markdown_report(r: ResumeAnalysis) -> str:
    lines = [
        "# ATS Resume Report",
        f"\n**Overall ATS score: {r.overall_score}/100 ({score_label(r.overall_score)})**\n",
        r.summary,
        "\n## Category scores",
    ]
    for name, value in r.category_scores.model_dump().items():
        lines.append(f"- {name.replace('_', ' ').title()}: {value}/100")
    lines += ["\n## Strengths"] + [f"- {s}" for s in r.strengths]
    lines += ["\n## Weaknesses"] + [f"- {s}" for s in r.weaknesses]
    lines += ["\n## Keywords found", ", ".join(r.found_keywords) or "None"]
    lines += ["\n## Missing keywords", ", ".join(r.missing_keywords) or "None"]
    lines.append("\n## Improvements")
    for i in r.improvements:
        lines.append(f"- **[{i.priority}] {i.area}**: {i.issue} -> {i.fix}")
    lines.append("\n## Bullet rewrites")
    for b in r.bullet_rewrites:
        lines.append(f"- Before: {b.original}\n  After: {b.improved}")
    return "\n".join(lines)


def render_results(r: ResumeAnalysis) -> None:
    st.divider()
    col1, col2 = st.columns([1, 2])
    with col1:
        st.metric("Overall ATS score", f"{r.overall_score}/100", score_label(r.overall_score), delta_color="off")
        st.progress(r.overall_score / 100)
    with col2:
        st.subheader("Summary")
        st.write(r.summary)

    st.subheader("Score breakdown")
    cats = r.category_scores.model_dump()
    cols = st.columns(len(cats))
    for col, (name, value) in zip(cols, cats.items()):
        with col:
            st.metric(name.replace("_", " ").title(), f"{value}")
            st.progress(value / 100)

    left, right = st.columns(2)
    with left:
        st.subheader("Strengths")
        for s in r.strengths:
            st.markdown(f"- {s}")
    with right:
        st.subheader("Weaknesses")
        for s in r.weaknesses:
            st.markdown(f"- {s}")

    kl, kr = st.columns(2)
    with kl:
        st.subheader("Keywords found")
        st.write(", ".join(f"`{k}`" for k in r.found_keywords) or "None detected")
    with kr:
        st.subheader("Missing keywords")
        st.write(", ".join(f"`{k}`" for k in r.missing_keywords) or "None, nice!")

    st.subheader("Recommended improvements")
    order = {"high": 0, "medium": 1, "low": 2}
    icons = {"high": "🔴", "medium": "🟠", "low": "🟢"}
    for imp in sorted(r.improvements, key=lambda i: order.get(i.priority.lower(), 3)):
        icon = icons.get(imp.priority.lower(), "⚪")
        with st.expander(f"{icon} {imp.priority} · {imp.area}"):
            st.markdown(f"**Issue:** {imp.issue}")
            st.markdown(f"**Fix:** {imp.fix}")

    if r.bullet_rewrites:
        st.subheader("Bullet point rewrites")
        for b in r.bullet_rewrites:
            st.markdown(f"**Before:** {b.original}")
            st.markdown(f"**After:** {b.improved}")
            st.write("")

    st.download_button(
        "Download report (.md)",
        data=build_markdown_report(r),
        file_name="ats_report.md",
        mime="text/markdown",
    )


def get_api_key(sidebar_value: str) -> str:
    if sidebar_value.strip():
        return sidebar_value.strip()
    try:
        if "GEMINI_API_KEY" in st.secrets:
            return str(st.secrets["GEMINI_API_KEY"]).strip()
    except Exception:
        pass  # no secrets file locally
    return os.getenv("GEMINI_API_KEY", "").strip()


# --------------------------------------------------------------------------- #
# App
# --------------------------------------------------------------------------- #
def main() -> None:
    st.set_page_config(page_title="AI Resume ATS Checker", page_icon="📄", layout="wide")
    st.title("📄 AI Resume ATS Checker")
    st.caption("Upload your resume, get an ATS score and concrete ways to improve it.")

    with st.sidebar:
        st.header("Settings")
        key_input = st.text_input(
            "Gemini API key",
            type="password",
            help="Get a free key at https://aistudio.google.com/apikey. "
            "On a deployed app, store it in Streamlit secrets instead.",
        )
        model = st.text_input("Gemini model", value=DEFAULT_MODEL)
        st.markdown("---")
        st.caption(
            "Your resume is sent to Google's Gemini API for analysis and is not stored by this app. "
            "Remove personal details if you prefer."
        )

    api_key = get_api_key(key_input)

    uploaded = st.file_uploader("Upload your resume", type=["pdf", "docx", "txt"])
    job_description = st.text_area(
        "Job description (optional, but gives a much more accurate score)",
        height=160,
        placeholder="Paste the job posting here to check keyword match...",
    )

    if st.button("Analyze resume", type="primary", disabled=uploaded is None):
        if not api_key:
            st.error("Please add your Gemini API key in the sidebar (or in Streamlit secrets).")
            st.stop()
        if uploaded.size > MAX_FILE_MB * 1024 * 1024:
            st.error(f"File is too large. Maximum size is {MAX_FILE_MB} MB.")
            st.stop()

        try:
            with st.spinner("Reading your resume..."):
                text = extract_resume_text(uploaded.name, uploaded.getvalue())
        except ValueError as exc:
            st.error(str(exc))
            st.stop()

        try:
            with st.spinner("Analyzing with Gemini..."):
                st.session_state["result"] = analyze_resume(api_key, text, job_description, model.strip() or DEFAULT_MODEL)
        except Exception as exc:
            st.error(friendly_error(exc))
            st.stop()

    if "result" in st.session_state:
        render_results(st.session_state["result"])


if __name__ == "__main__":
    main()
